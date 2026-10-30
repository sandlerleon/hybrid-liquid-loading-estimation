# -*- coding: utf-8 -*-
"""Revision 2 of FLOWMEASINST-D-26-00971: response to the reviewers, cover letter, pasteable replies for the Elsevier review
platform, title page with updated metrics, TIFF figures and the submission folder. Run build_fmi_v6.py first.

    python build_fmi_revision2_letters.py
"""
import copy
import os
import re
import shutil

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, ".."))
SRCD = r"C:\Users\Leon\Downloads\FMI Manuscript"
OUTD = os.path.join(SRCD, "revision2")
os.makedirs(os.path.join(OUTD, "figures"), exist_ok=True)

DATE = "29 October 2026"
MS = "FLOWMEASINST-D-26-00971"
TITLE = "Hybrid Physics–Diagnostic Estimation of Liquid Loading for Wet-Gas Flowmeter Correction: A Model-Based Uncertainty Framework"
MAN = os.path.join(REPO, "manuscript", "FMI_Hybrid_Liquid_Loading_Estimation_v6.docx")


def fix_zoom(doc):
    """python-docx's default settings.xml has <w:zoom> without the required w:percent attribute."""
    z = doc.settings.element.find(qn("w:zoom"))
    if z is not None and z.get(qn("w:percent")) is None:
        z.set(qn("w:percent"), "100")

R2_COMMENT = ("Despite the revisions, the manuscript remains poorly structured and presented. The text often resembles a fragment of program "
              "code or a technical development log rather than a properly written scientific paper. Look at your equations!!! The excessive "
              "use of fragmented explanations, implementation-oriented details, and abrupt transitions significantly impairs readability and "
              "obscures the scientific narrative. The manuscript requires substantial editorial restructuring to ensure a coherent "
              "presentation of the research problem, methodology, results, and scientific contribution. The authors should rewrite the text "
              "in a clear, concise, and academically appropriate style, rather than merely correcting individual formatting issues. The "
              "overall presentation must meet the standards expected of a scientific publication.")

CHANGES = [
    ("Equations.",
     "All seventeen equations are now display equations typeset in the Office Math equation editor, with proper fractions, radicals, "
     "subscripts and superscripts and the equation number at the right margin. In the previous version they were plain text strings such "
     "as “Q_ind = Q_G · φ(X_LM, ρ_G/ρ_L, Fr_G)”, which is what the reviewer saw. Symbols in the running text are set the same way: "
     "italic variables and upright labels, so that X_LM, u₁ and ρ₁₂ no longer appear with underscores. The equations are gathered in one "
     "section (Section 3, Mathematical model), which opens by stating the order in which they are applied. Each group is introduced by a "
     "sentence and followed by a paragraph that explains its origin and assumptions."),
    ("Program-code content removed.",
     "Algorithm 1, which was written as pseudo-code, is removed. The decision procedure is now described in prose (Section 2.4), and the "
     "order of the calculation is given in the opening paragraph of Section 3. Script and function names no longer appear in the main "
     "text; they are confined to the Data availability statement and to the provenance table in Appendix A (Table A1)."),
    ("Development-log content removed.",
     "Statements about what had been added or left unchanged in the revision are removed. Examples are “(new)”, “the new analyses are in "
     "Sections 5.9–5.12” and “Tables 3, 5 and 6 and Figs 2–8 are the original results and are unchanged”. The two Monte Carlo samples "
     "remain, because they are a property of the method; they are now described once, in Section 3.2, as part of the design of the study."),
    ("Structure.",
     "The paper now runs in a single sequence: research problem, estimation architecture, mathematical model, computational design, "
     "results, discussion and conclusions. The four-part introduction with its numbered list of contributions is now one continuous "
     "introduction. It moves from the literature to the gap and then to three research questions, whose answers are stated at the end of "
     "the introduction and again in the conclusions. Third-level headings, bold run-in labels and bulleted lists, including the bulleted "
     "Limitations, are replaced by connected paragraphs. Each Results subsection opens with the question it answers and closes with its "
     "interpretation. The validation design, previously in the Methods, is now in the Discussion (Section 5.5), next to the limitations "
     "it addresses."),
    ("Provenance and verification.",
     "A new Section 3.1 states the origin of each equation (established, adapted from the literature, or specific to this study) and its "
     "assumptions. Appendix A gives this equation by equation (Table A1) and reports the executable checks of the implementation "
     "(Table A2). The implementation details that the reviewer found intrusive in the text are kept, but in this one place."),
    ("Figures, references and nomenclature.",
     "The figure labels now match the text. The internal acronym “SWCL” and the shorthand “E1/E2” are removed from Figs. 2, 3 and 12. The "
     "title of Fig. 7 no longer uses “accuracy”, which the paper reserves for measured quantities. Fig. 1 now shows the consistency index "
     "with the covariance term of Eq. (11) and uses the symbols of the text. The figures were regenerated from the deposited code with "
     "the fixed seeds, and their data are unchanged. References are numbered in order of citation. The nomenclature, with a units "
     "column, is placed after the Conclusions."),
]

MAPPING = [
    ("Abstract", "Abstract (rewritten; same results)"),
    ("1.1–1.4 Introduction (four subsections, numbered list of contributions)", "1 Introduction (continuous; research questions, novelty, status of the results)"),
    ("2 Problem formulation", "2.1 The measurement problem; Eqs. (1)–(4) in Section 3"),
    ("3.1 Architecture", "2 (opening paragraph) and Fig. 1"),
    ("3.2 Thermodynamic estimator", "2.2 Thermodynamic estimator"),
    ("3.3, 3.3.1 Diagnostic estimator and its specification", "2.3 Diagnostic estimator (Table 1)"),
    ("3.4 Reference-based calibration", "2.4 Calibration, fusion and the correction decision"),
    ("3.5, 3.5.1, 3.5.2 Fusion, damping, propagation", "3 Mathematical model (Eqs. (8)–(16) and their explanation)"),
    ("3.6 Algorithm 1 (pseudo-code)", "Removed; described in prose in 2.4 and at the opening of 3"),
    ("–", "3.1 Origin and assumptions of the equations; Appendix A, Table A1"),
    ("4.1, 4.2, 4.5 Forward model, sampling, additional experiments", "3.2 Monte Carlo propagation and operating envelope (Table 2)"),
    ("4.3 Metrics and terminology", "3.3 Metrics, terminology and verification (Eq. (17); Table A2)"),
    ("4.4 Validation design and falsification criteria", "5.5 Proposed validation"),
    ("5.1, 5.2", "4.1, 4.2"),
    ("5.3 Consistency index; 5.4 Threshold selection", "4.3 The consistency index and the acceptance threshold"),
    ("5.5–5.12", "4.4–4.11, in the same order"),
    ("6.1, 6.2, 6.3", "5.1, 5.2, 5.3"),
    ("6.4 Limitations (bulleted)", "5.4 Limitations (paragraphs)"),
    ("7 Conclusions", "6 Conclusions"),
    ("Nomenclature (before the Introduction)", "Nomenclature (after the Conclusions, with units)"),
]

# ============================================================================================ manuscript metrics
man = Document(MAN)
ps = [p.text for p in man.paragraphs]
i0 = next(k for k, t in enumerate(ps) if t.strip() == "1. Introduction")
i1 = next(k for k, t in enumerate(ps) if t.strip() == "Nomenclature")
words = sum(len(t.split()) for t in ps[i0:i1])
iabs = next(k for k, t in enumerate(ps) if t.strip() == "Abstract")
abs_words = len(ps[iabs + 1].split())
caps = [t for t in ps if re.match(r"^(Fig\.|Table) [0-9A]+\. ", t.strip())]
nfig = len([t for t in caps if t.startswith("Fig.")])
ntab = len([t for t in caps if re.match(r"^Table \d+\.", t)])
ntabA = len([t for t in caps if re.match(r"^Table A\d+\.", t)])
nref = len([t for t in ps if re.match(r"^\[\d+\] ", t)])
print("metrics: abstract %d, main text %d, figures %d, tables %d + %d, references %d" % (abs_words, words, nfig, ntab, ntabA, nref))

# ============================================================================================ response document
d = Document()
st = d.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(6)
for s_ in d.sections:
    s_.left_margin = s_.right_margin = Inches(1.1)
    s_.top_margin = s_.bottom_margin = Inches(1.0)


def para(text="", bold=False, italic=False, size=None, indent=None, after=None, before=None):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    return p


def labelled(label, text, indent=None):
    p = d.add_paragraph()
    r = p.add_run(label + " ")
    r.bold = True
    p.add_run(text)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    return p


para("Response to Reviewers — Revision 2", bold=True, size=15, after=2)
para("Manuscript " + MS, bold=True, after=2)
para(TITLE, italic=True, after=2)
para("Leon Sandler, Independent Researcher, Northbrook, Illinois, USA", after=2)
para(DATE, after=10)
para("Dear Dr Kutin and Reviewers,")
para("Thank you for the second round of review. Reviewer 1 found that all comments had been addressed. Reviewer 2 found that the "
     "manuscript was still poorly structured and presented, and asked for a rewrite rather than local corrections. I agree with that "
     "assessment, and the manuscript has been rewritten from beginning to end. No calculation, result or conclusion has changed: every number "
     "in the revised text is the number of the previous version, and the deposited code reproduces it. The revised manuscript is "
     "FMI_Hybrid_Liquid_Loading_Estimation_v6.docx; section numbers below refer to it.")

para("Reviewer 1", bold=True, size=13, before=12)
labelled("Comment.", "All comments have been addressed satisfactorily. No further comments.", 0.25)
labelled("Response.", "Thank you for the review and for confirming the previous revision. The present revision restructures and rewrites "
         "the presentation in response to Reviewer 2. The analyses, results and conclusions that you reviewed are unchanged, and the "
         "points you raised (the framework as the principal contribution, the conditions under which fusion helps, precision rather than "
         "specificity, the recalculation after damping, the uncertainty treatment, the minimum signal set and the units) remain in the "
         "text, now in Sections 1, 4.5, 4.3, 3, 2.3, 5.2 and 3.2 respectively.", 0.25)

para("Reviewer 2", bold=True, size=13, before=12)
labelled("Comment.", R2_COMMENT, 0.25)
labelled("Response.", "Thank you for this comment, which is justified. The previous revision answered the ten scientific points of the first "
         "review by adding analyses where they were needed, and the result read like the record of a development rather than a paper. The "
         "equations were the clearest symptom. The manuscript has now been rewritten as a whole: its structure, its prose and its "
         "typesetting are new, and its content and numbers are unchanged. The changes are as follows.", 0.25)
for lab, txt in CHANGES:
    labelled(lab, txt, 0.5)
para("Correspondence between the previous and the present version", bold=True, before=6, indent=0.25)
t = d.add_table(rows=1, cols=2)
t.style = "Light Grid Accent 1"
t.rows[0].cells[0].text = "Previous version (v5)"
t.rows[0].cells[1].text = "Present version (v6)"
for a, b in MAPPING:
    row = t.add_row().cells
    row[0].text, row[1].text = a, b
for row in t.rows:
    for c in row.cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)
para("")
labelled("Unchanged.", "The equation numbers (1)–(17), the table numbers 1–10 and the figure numbers 1–12 are the same as in the previous "
         "version, so that the responses of the first round still point to the right objects.", 0.25)
labelled("Marked copy.", "Because the text was rewritten throughout, a change-marked copy would mark every paragraph and would not help the "
         "reader. The table above is provided in its place, and the previous version remains available for comparison.", 0.25)
para("Closing", bold=True, size=13, before=10)
para("I am grateful to Reviewer 2 for insisting on this. The paper now presents the problem, the method, the results and their limits "
     "as one argument, and its equations can be read as equations. I would be glad to make any further changes the editor or the reviewers "
     "consider necessary.")
para("Leon Sandler")
fix_zoom(d)
d.save(os.path.join(OUTD, "FMI_Response_to_Reviewers_R2.docx"))

# ============================================================================================ pasteable replies
rep1 = ("Thank you for the review and for confirming the previous revision. The present revision restructures and rewrites the presentation "
        "in response to Reviewer 2; the analyses, results and conclusions you reviewed are unchanged, and every point you raised remains "
        "in the text. The full response is in the uploaded file FMI_Response_to_Reviewers_R2.docx.")
rep2 = ["Thank you for this comment, which is justified. The manuscript has been rewritten from beginning to end rather than corrected "
        "locally; its structure, prose and typesetting are new, and no calculation, result or conclusion has changed. The full response, "
        "with a table of correspondence between the two versions, is in the uploaded file FMI_Response_to_Reviewers_R2.docx. In brief:", ""]
for i, (lab, txt) in enumerate(CHANGES, 1):
    rep2 += ["%d. %s %s" % (i, lab, txt), ""]
rep2 += ["Equation, table and figure numbers are unchanged. Because the text was rewritten throughout, a change-marked copy would mark every "
         "paragraph; the correspondence table in the response file is provided in its place."]
open(os.path.join(OUTD, "reply_reviewer1.txt"), "w", encoding="utf-8").write(rep1 + "\n")
open(os.path.join(OUTD, "reply_reviewer2.txt"), "w", encoding="utf-8").write("\n".join(rep2) + "\n")

# ============================================================================================ cover letter
cl = Document(os.path.join(SRCD, "FMI_Cover_Letter_Sandler.docx"))
cp = cl.paragraphs
tm = dict(name=cp[0], line=cp[1], date=cp[4], to=cp[5], body=cp[9], sign=cp[20])
holders = {k: Paragraph(copy.deepcopy(v._p), cl._body) for k, v in tm.items()}
body = cl.element.body
for el in list(body):
    if not el.tag.endswith("}sectPr"):
        body.remove(el)


def cmake(kind, text, bold=None):
    tmpl = holders[kind]
    p = copy.deepcopy(tmpl._p)
    for r in list(p):
        if r.tag == qn("w:r"):
            p.remove(r)
    body.find(qn("w:sectPr")).addprevious(p)
    par = Paragraph(p, cl._body)
    r = copy.deepcopy(tmpl.runs[0]._r)
    for x in list(r):
        if x.tag in (qn("w:t"), qn("w:br"), qn("w:tab")):
            r.remove(x)
    p.append(r)
    rr = Run(r, par)
    rr.text = text
    if bold is not None:
        rr.bold = bold
    return par


cmake("name", "Leon Sandler", True)
for t_ in ("Independent Researcher", "Northbrook, IL 60062, USA", "sandler.leon@gmail.com"):
    cmake("line", t_)
cmake("date", DATE)
cmake("to", "Dr Joze Kutin")
cmake("to", "Senior Associate Editor, Flow Measurement and Instrumentation")
cmake("to", "Elsevier")
cmake("body", "Re: Revision 2 of manuscript %s — “%s”" % (MS, TITLE), True)
cmake("body", "Dear Dr Kutin,")
cmake("body", "Thank you for the second round of review. Reviewer 1 found that all comments had been addressed. Reviewer 2 found the "
              "manuscript still poorly structured and presented, with text that read like program code or a development log and equations "
              "written as plain text, and asked for a rewrite rather than local corrections.")
cmake("body", "I agree, and the manuscript has been rewritten from beginning to end. Its equations are now typeset display equations; the "
              "pseudo-code algorithm, the revision notes and the implementation details have been removed from the text; and the paper "
              "follows a single sequence from the research problem through the architecture, the mathematical model and the computational "
              "design to the results, discussion and conclusions. A new section states the origin and assumptions of each equation, and an "
              "appendix reports the provenance of the equations and the verification of their implementation. Figure labels have been "
              "harmonized with the text.")
cmake("body", "No calculation, result or conclusion has changed. Every number in the revised manuscript is that of the previous version, "
              "and the equation, table and figure numbers are the same. The response to the reviewers includes a table of correspondence "
              "between the two versions; because the text was rewritten throughout, it is provided in place of a change-marked copy.")
cmake("body", "The code, including the scripts that build the manuscript, remains openly available at "
              "https://github.com/sandlerleon/hybrid-liquid-loading-estimation (archived at https://doi.org/10.5281/zenodo.22760863).")
cmake("body", "The manuscript is original, is not under consideration elsewhere, and has not been published previously except as the ChemRxiv "
              "preprint of its first version (https://doi.org/10.26434/chemrxiv.15009070/v1). I am the sole author, I have no competing "
              "interests, and the work received no external funding. The use of a generative AI tool, including in the restructuring of "
              "the present revision, is disclosed in the manuscript.")
cmake("body", "Thank you for your consideration.")
cmake("sign", "Sincerely,")
cmake("sign", "Leon Sandler", True)
cmake("sign", "Independent Researcher, Northbrook, IL, USA")
cmake("sign", "sandler.leon@gmail.com")
fix_zoom(cl)
cl.save(os.path.join(OUTD, "FMI_Revision2_Cover_Letter.docx"))

# ============================================================================================ title page and highlights
tp = Document(os.path.join(SRCD, "revision1", "FMI_Title_Page.docx"))
for p in tp.paragraphs:
    t_ = p.text
    new = None
    if t_.startswith("Abstract: "):
        new = "Abstract: %d words (unstructured, single paragraph)" % abs_words
    elif t_.startswith("Main text: "):
        new = "Main text: approximately %s words, excluding nomenclature, back matter, appendix and references" % format(int(round(words, -1)), ",")
    elif t_.startswith("Figures: "):
        new = "Figures: %d (all original; generated by the deposited code)" % nfig
    elif t_.startswith("Tables: "):
        new = "Tables: %d, and %d in Appendix A" % (ntab, ntabA)
    elif t_.startswith("References: "):
        new = "References: %d" % nref
    if new:
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].text = new
fix_zoom(tp)
tp.save(os.path.join(OUTD, "FMI_Title_Page.docx"))
hl = Document(os.path.join(SRCD, "revision1", "FMI_Highlights.docx"))
fix_zoom(hl)
hl.save(os.path.join(OUTD, "FMI_Highlights.docx"))

# ============================================================================================ manuscript and figures
for suf in ("", "_no_author_block"):
    shutil.copy2(os.path.join(REPO, "manuscript", "FMI_Hybrid_Liquid_Loading_Estimation_v6%s.docx" % suf), OUTD)
FIGS = [("Fig1_method", "figG_method"), ("Fig2_ablation", "figB_ablation"), ("Fig3_distribution", "figE_distribution"),
        ("Fig4_agreement", "figA_estimator_agreement"), ("Fig5_covariance", "figH_covariance"), ("Fig6_designspace", "figF_designspace"),
        ("Fig7_sensitivity", "figC_sensitivity"), ("Fig8_envelope", "figD_envelope"), ("Fig9_detector", "figK_detector"),
        ("Fig10_robustness", "figJ_robustness"), ("Fig11_modelform", "figI_modelform"), ("Fig12_deployment", "fig6_deployment_topology")]
for out, src in FIGS:
    im = Image.open(os.path.join(REPO, "figures", src + ".png")).convert("RGB")
    im.save(os.path.join(OUTD, "figures", out + ".tif"), compression="tiff_lzw", dpi=(300, 300))
open(os.path.join(OUTD, "MANIFEST.txt"), "w", encoding="utf-8").write(
    "%s - Revision 2 - prepared %s\n\n"
    "Upload in Editorial Manager:\n"
    "  FMI_Hybrid_Liquid_Loading_Estimation_v6.docx                 revised manuscript (clean; rewritten throughout)\n"
    "  FMI_Hybrid_Liquid_Loading_Estimation_v6_no_author_block.docx anonymized copy, if requested\n"
    "  FMI_Response_to_Reviewers_R2.docx                            response, with the v5 -> v6 correspondence table\n"
    "  FMI_Revision2_Cover_Letter.docx                              cover letter\n"
    "  FMI_Highlights.docx                                          unchanged\n"
    "  FMI_Title_Page.docx                                          metrics updated\n"
    "  figures/Fig1-Fig12 (.tif, 300 dpi)                           Figs 1, 2, 3, 7, 12 relabeled; data unchanged\n\n"
    "Review platform (one reply per comment):\n"
    "  reply_reviewer1.txt   paste under Reviewer 1 - Revision 1\n"
    "  reply_reviewer2.txt   paste under Reviewer 2 - Revision 1\n\n"
    "No change-marked copy: the text was rewritten throughout (explained in the response).\n"
    "Manuscript metrics: abstract %d words; main text about %s words; %d figures; %d tables + %d appendix tables; %d references.\n"
    % (MS, DATE, abs_words, format(int(round(words, -1)), ","), nfig, ntab, ntabA, nref))
print("saved to", OUTD)
