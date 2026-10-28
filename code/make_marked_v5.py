# -*- coding: utf-8 -*-
"""Change-marked copy of manuscript v5: every paragraph or table that differs from the submitted v4 is highlighted.

    python make_marked_v5.py
"""
import os

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, ".."))
V4 = r"C:\Users\Leon\Downloads\FMI Manuscript\FMI_Hybrid_Liquid_Loading_Estimation_v4.docx"
V5 = os.path.join(REPO, "manuscript", "FMI_Hybrid_Liquid_Loading_Estimation_v5.docx")
OUT = os.path.join(REPO, "manuscript", "FMI_Hybrid_Liquid_Loading_Estimation_v5_marked.docx")

a, b = Document(V4), Document(V5)
old_pars = {p.text.strip() for p in a.paragraphs}
old_tabs = {tuple(c.text for r in t.rows for c in r.cells) for t in a.tables}


def hl(par):
    for r in par.runs:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW


n_par = n_tab = 0
for p in b.paragraphs:
    t = p.text.strip()
    if t and t not in old_pars and not p._p.findall(".//" + qn("a:blip")):
        hl(p)
        n_par += 1
    elif p._p.findall(".//" + qn("a:blip")):
        pass
for t in b.tables:
    sig = tuple(c.text for r in t.rows for c in r.cells)
    if sig not in old_tabs:
        n_tab += 1
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    hl(p)
b.save(OUT)
tot = sum(1 for p in b.paragraphs if p.text.strip())
print("highlighted %d paragraphs of %d and %d tables ; saved %s" % (n_par, tot, n_tab, OUT))
