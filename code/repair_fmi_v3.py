# -*- coding: utf-8 -*-
"""Repair the externally copyedited FMI manuscript.

The copyedit made many legitimate US-English improvements and left every number,
DOI, equation number and Unicode symbol intact. It also introduced a corrupted
token, several reversed or garbled claims, and altered four published reference
titles. This restores the damage and keeps the rest.

    python repair_fmi_v3.py  ->  FMI_Hybrid_Liquid_Loading_Estimation_v4.docx
"""
import io, os, re, shutil
from docx import Document

D = r"C:\Users\Leon\Downloads\FMI Manuscript"
SRC = os.path.join(D, "FMI_Hybrid_Liquid_Loading_Estimation_v3.docx")
DST = os.path.join(D, "FMI_Hybrid_Liquid_Loading_Estimation_v4.docx")
shutil.copyfile(SRC, DST)
doc = Document(DST)
log = []


def set_text(par, t):
    rs = par.runs
    if not rs:
        par.add_run(t); return
    rs[0].text = t
    for r in rs[1:]:
        r.text = ""


def targets():
    out = list(doc.paragraphs)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                out.extend(c.paragraphs)
    return out


def fix(old, new, label, all_occurrences=False):
    n = 0
    for par in targets():
        if old in par.text:
            set_text(par, par.text.replace(old, new))
            n += 1
            if not all_occurrences:
                break
    log.append(("OK  " if n else "!!  MISS ") + label + (f"  [{n}]" if n else ""))
    return n


# ---------------------------------------------- corrupted token and lost meaning
fix("UnderrepresentUnderative envelope conditions, the condensate–gas ratios of 50 and 100 "
    "bbl/MMscf are greater than 7.9% and 15.7%, respectively",
    "At representative envelope conditions, condensate–gas ratios of 50 and 100 bbl/MMscf give "
    "over-readings of 7.9% and 15.7%",
    "corrupted token 'UnderrepresentUnderative' + 'are greater than' destroyed the claim")
fix("which is necessary to check everything downstream",
    "which is a necessary check on everything downstream", "reversed meaning")

fix("At most production metering stations for which parameters are not measured, the correction "
    "cannot be closed to the meter signal alone.",
    "At most production metering stations that parameter is not measured, so the correction "
    "cannot be closed from the meter signal alone.",
    "abstract: restrictive clause and 'closed to' both wrong")
fix("In this paper, liquid loading is treated as a latent variable to be estimated rather than "
    "as a quantity to be learned.",
    "This paper treats liquid loading as a latent variable to be estimated rather than treating "
    "the over-reading as a quantity to be learned.",
    "abstract: lost the contrast with the over-reading")

fix("X_LM is estimated from process and meter information already available at the station, the "
    "uncertainty of that estimate is quantified, and X_LM is propagated through Eq. (2) to "
    "determine the defensible uncertainty in the corrected gas rate.",
    "Estimate X_LM from process and meter information already available at the station, quantify "
    "the uncertainty of that estimate, and propagate it through Eq. (2) to a defensible "
    "uncertainty on the corrected gas rate.",
    "problem statement: it is the uncertainty that is propagated, not X_LM")

fix("so we assume independence when ρ₁₂ > 0 makes the denominator too large",
    "so assuming independence when ρ₁₂ > 0 makes the denominator too large",
    "ungrammatical, and asserts the opposite of the recommendation")

fix("Fusing while the thermodynamic estimator declares that only its flash uncertainty is "
    "±5.45%, which is worse than that of the diagnostic estimator alone",
    "Fusing while the thermodynamic estimator declares only its flash uncertainty gives ±5.45%, "
    "worse than the diagnostic estimator alone",
    "garbled: ±5.45% is the fused band, not the flash uncertainty")

fix("Z_max = 3 retains almost all the data when the cases in which the estimator cannot handle "
    "the data are removed.",
    "Z_max = 3 retains almost all data while removing the cases the estimator cannot handle.",
    "circular and garbled")

fix("Liquid injection upstream of the flash-equilibrium point is delayed in a subset of runs to "
    "create the carryover condition",
    "Deliberate liquid injection upstream of the flash-equilibrium point, in a subset of runs, to "
    "create the carryover condition", "'delayed' should be 'deliberate'")
fix("and either a chromatograph or a gas of known composition were used.",
    "and either a chromatograph or a gas of known composition.",
    "past tense asserts an experiment that has not been performed")
fix("Falsification criteria The framework is refuted",
    "Falsification criteria. The framework is refuted", "deleted full stop")

fix("so that the study can be reproduced so that a reader can see",
    "so that the study can be reproduced and so that a reader can see", "deleted conjunction")

fix("with the transmitter, chromatograph and meter diagnostic and the corrected rate, its "
    "uncertainty, the fused liquid loading and the status flag being read.",
    "reading the transmitters, chromatograph and meter diagnostics and publishing the corrected "
    "rate, its uncertainty, the fused liquid loading and the status flag.",
    "destroyed the read/publish distinction, which is the point of the sentence")

fix("which is the condition under which an inlet separator and dehydration train are produced",
    "which is the condition an inlet separator and dehydration train produce",
    "inverted: the train produces the condition, not the reverse")

fix("which is why Section 4.4 performs that measurement part of the validation protocol",
    "which is why Section 4.4 makes that measurement part of the validation protocol",
    "'performs' should be 'makes'")

fix("a weaker result than direct learning can be achieved when training data are plentiful",
    "a weaker result than direct learning can achieve where training data are plentiful",
    "garbled")

fix("in 94% of those the tests reject", "in 94% of those the test rejects", "subject/verb")

fix("available where downstream tapping occurs", "available where a downstream tapping exists",
    "a tapping is a fitting; it does not 'occur'")
fix("the permanent pressure loss occurs because the film alters the recovery",
    "the permanent pressure loss responds because the film alters the recovery",
    "the loss responds to the film; it does not occur because of it")

fix("Declaration of competing interests", "Declaration of Competing Interest",
    "Elsevier's required section heading")

# ------------------------------------------------- published reference titles
fix("A new hybrid physics–ML model for predicting overreading in differential pressure",
    "A new hybrid physics–ML model for predicting over-reading in differential pressure",
    "ref: Alkhurayef 2026a title is 'over-reading'")
fix("enhancing generalization in digital measurement systems",
    "enhancing generalisation in digital measurement systems",
    "ref: Hosseini 2026a title uses British spelling")
fix("in an orifice plate meter for gas‒liquid systems",
    "in an orifice plate meter for gas-liquid systems",
    "ref: Nasr 2026 — figure dash substituted for hyphen")
fix("wet gas flow with ultralow liquid loading through DP meters",
    "wet gas flow with ultra-low liquid loading through DP meters",
    "ref: Zheng 2019 title is 'ultra-low'")
fix("American Gas Association, Report No. 3/API MPMS Chapter 14.3",
    "American Gas Association, Report No. 3 / API MPMS Chapter 14.3", "standard designation spacing")
fix("American Gas Association, Report No. 11/API MPMS Chapter 14.9",
    "American Gas Association, Report No. 11 / API MPMS Chapter 14.9", "standard designation spacing")

# ------------------------------------------------------ global consistency
# 'over-reading' is the spelling used by the cited literature, including the
# reference titles restored above; keep the document internally consistent.
n = 0
for par in targets():
    t = par.text
    if "overreading" in t or "Overreading" in t:
        set_text(par, t.replace("overreading", "over-reading").replace("Overreading", "Over-reading"))
        n += 1
log.append(f"SWEEP overreading -> over-reading: {n} paragraph(s)")

n = 0
for par in targets():
    t = par.text
    if "standing pseudocritical" in t or "Standing pseudocritical" in t:
        set_text(par, t.replace("standing pseudocritical", "Standing pseudo-critical")
                       .replace("Standing pseudocritical", "Standing pseudo-critical"))
        n += 1
    elif "Standing pseudocriticals" in t or "standing pseudocriticals" in t:
        set_text(par, t.replace("Standing pseudocriticals", "Standing pseudo-criticals")
                       .replace("standing pseudocriticals", "Standing pseudo-criticals"))
        n += 1
log.append(f"SWEEP Standing pseudo-critical (proper name): {n} paragraph(s)")

doc.save(DST)
io.open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "_repair_fmi_log.txt"),
        "w", encoding="utf-8").write("\n".join(log))
print("\n".join(log))
print("\nsaved", DST)
