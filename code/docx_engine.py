# -*- coding: utf-8 -*-
"""Document engine for the manuscript: reproduces the layout of the author's Desalination and Water Treatment
manuscript (US Letter, Calibri 10.5 pt, numbered display equations in Office Math, Light Grid tables with captions
below, numbered references in order of first citation).

Inline markup used in the text blocks:
    $...$          inline mathematics: Latin letters italic, \\greek names, _{sub} and ^{sup}, \\mathrm{upright},
                   \\tilde{w}, \\dot{m}
    [@a;@b]        citation keys, numbered in order of first appearance
"""
import copy
import os
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, Emu

import omml

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "manuscript_template.docx")
TEXT_WIDTH_TW = 9792            # 12240 - 2 x 1224 twips

GREEK = dict(omml.GREEK)
GREEK.update({"Delta": "Δ", "Sigma": "Σ", "theta": "θ", "varepsilon": "ε", "epsilon": "ε", "phi": "φ", "varphi": "φ",
              "rho": "ρ", "sigma": "σ", "Upsilon": "Υ", "le": "≤", "ge": "≥", "pm": "±", "times": "×", "cdot": "·",
              "minus": "−", "prime": "′", "approx": "≈", "infty": "∞", "partial": "∂", "sqrt": "√"})
UPRIGHT_GREEK = set("ΔΣΥΓΛΩΦ")
FUNCS = {"ln", "max", "min", "var", "exp", "log", "d"}
ACCENT = {"tilde": "̃", "dot": "̇", "hat": "̂", "bar": "̅"}


# ------------------------------------------------------------------------------------------- run formatting
def _set_rpr(run, size, italic=False, bold=False, va=None):
    run.font.size = Pt(size / 2.0)
    run.italic = bool(italic)
    if bold:
        run.bold = True
    if va == "sub":
        run.font.subscript = True
    elif va == "sup":
        run.font.superscript = True


def _group(s, i):
    """Return (content, next index) of a {group} or a single character / command at s[i]."""
    if i >= len(s):
        return "", i
    if s[i] == "{":
        depth, j = 1, i + 1
        while j < len(s) and depth:
            if s[j] == "{":
                depth += 1
            elif s[j] == "}":
                depth -= 1
            j += 1
        return s[i + 1:j - 1], j
    if s[i] == "\\":
        m = re.match(r"\\([A-Za-z]+|.)", s[i:])
        return m.group(0), i + len(m.group(0))
    return s[i], i + 1


def _math_tokens(s, upright=False):
    """Yield (text, italic, va) pieces for an inline-math string."""
    out = []
    i = 0
    while i < len(s):
        c = s[i]
        if c == "\\":
            m = re.match(r"\\([A-Za-z]+|.)", s[i:])
            name = m.group(1)
            i += len(m.group(0))
            if name in ("mathrm", "rm", "text"):
                g, i = _group(s, i)
                out += [(t, False, va) for t, _, va in _math_tokens(g, upright=True)]
            elif name in ACCENT:
                g, i = _group(s, i)
                sub = _math_tokens(g, upright)
                if sub:
                    t, it, va = sub[0]
                    sub[0] = (t[0] + ACCENT[name] + t[1:], it, va)
                out += sub
            elif name in FUNCS:
                out.append((name, False, None))
            elif name in GREEK:
                ch = GREEK[name]
                out.append((ch, (not upright) and ch.isalpha() and ch not in UPRIGHT_GREEK, None))
            elif name in (",", ";", " "):
                out.append((" ", False, None))
            else:
                out.append((name, False, None))
        elif c in "_^":
            g, i = _group(s, i + 1)
            va = "sub" if c == "_" else "sup"
            for t, it, _ in _math_tokens(g, upright=True):
                out.append((t, it, va))
        elif c == "{" or c == "}":
            i += 1
        elif c.isalpha():
            out.append((c, not upright, None)); i += 1
        elif c == "-":
            out.append(("−", False, None)); i += 1
        else:
            out.append((c, False, None)); i += 1
    return out


def add_runs(par, text, size=21, italic=False, bold=False):
    """Add marked-up text to a paragraph. Inside captions (italic=True) mathematics stays italic throughout."""
    parts = text.split("$")
    for k, seg in enumerate(parts):
        if not seg:
            continue
        if k % 2 == 0:
            r = par.add_run(seg)
            _set_rpr(r, size, italic, bold)
        else:
            pieces = _math_tokens(seg)
            merged = []
            for t, it, va in pieces:
                if merged and merged[-1][1] == it and merged[-1][2] == va:
                    merged[-1] = (merged[-1][0] + t, it, va)
                else:
                    merged.append((t, it, va))
            for t, it, va in merged:
                r = par.add_run(t)
                _set_rpr(r, size, italic or it, bold, va)
    return par


# ------------------------------------------------------------------------------------------ paragraph kinds
def _spacing(par, before, after, line=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before / 20.0)
    pf.space_after = Pt(after / 20.0)
    if line:
        pf.line_spacing = line


class Doc:
    def __init__(self, citations):
        self.d = docx.Document(TEMPLATE)
        self.cit = citations

    def body_text(self, text):
        return self.cit.resolve(text)

    def title(self, text):
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, text, 30, bold=True)

    def centered(self, text, size=22):
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(p, 0, 20)
        add_runs(p, text, size)

    def blank(self):
        self.d.add_paragraph()

    def h1(self, text):
        p = self.d.add_paragraph(); _spacing(p, 300, 100)
        p.paragraph_format.keep_with_next = True
        add_runs(p, text, 26, bold=True)

    def h2(self, text):
        p = self.d.add_paragraph(); _spacing(p, 220, 100)
        p.paragraph_format.keep_with_next = True
        add_runs(p, text, 23, bold=True)

    def p(self, text, indent=True):
        par = self.d.add_paragraph(); _spacing(par, 0, 160)
        if indent:
            par.paragraph_format.first_line_indent = Emu(317 * 635)
        add_runs(par, self.body_text(text), 21)
        return par

    def eq(self, tex, number):
        par = self.d.add_paragraph()
        pPr = par._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        for pos, val in ((TEXT_WIDTH_TW // 2, "center"), (TEXT_WIDTH_TW, "right")):
            t = OxmlElement("w:tab"); t.set(qn("w:val"), val); t.set(qn("w:pos"), str(pos)); tabs.append(t)
        pPr.append(tabs)
        _spacing(par, 120, 120)
        par.paragraph_format.keep_together = True
        r = par.add_run(); r.add_tab()
        par._p.append(omml.tex_to_omml(tex))
        r2 = par.add_run(); r2.add_tab(); r2.add_text("(%d)" % number); r2.font.size = Pt(10.5)

    def fig(self, path, caption, width_in=6.4):
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(par, 200, 80); par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path, width=Inches(width_in))
        cap = self.d.add_paragraph(); _spacing(cap, 0, 200)
        add_runs(cap, self.body_text(caption), 18, italic=True)

    def table(self, caption, header, rows, widths=None, first_col_bold=True, size=17):
        n = len(header)
        widths = widths or [1.0] * n
        tot = float(sum(widths))
        tw = [int(TEXT_WIDTH_TW * w / tot) for w in widths]
        tw[-1] = TEXT_WIDTH_TW - sum(tw[:-1])
        t = self.d.add_table(rows=1 + len(rows), cols=n)
        tbl = t._tbl
        tblPr = tbl.tblPr
        st = OxmlElement("w:tblStyle"); st.set(qn("w:val"), "LightGridAccent1")
        old = tblPr.find(qn("w:tblStyle"))
        if old is not None:
            tblPr.remove(old)
        tblPr.insert(0, st)
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
        tblW.set(qn("w:w"), str(TEXT_WIDTH_TW)); tblW.set(qn("w:type"), "dxa")
        look = tblPr.find(qn("w:tblLook"))
        if look is None:
            look = OxmlElement("w:tblLook"); tblPr.append(look)
        # schema order: ... tblW, jc, tblCellSpacing, tblInd, tblBorders, shd, tblLayout, tblCellMar, tblLook
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
        mar = tblPr.find(qn("w:tblCellMar"))
        (mar if mar is not None else look).addprevious(lay)
        look.set(qn("w:val"), "04a0" if first_col_bold else "0420")
        grid = tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), tw):
            gc.set(qn("w:w"), str(w))
        for ri, row in enumerate([header] + rows):
            tr = t.rows[ri]
            if ri == 0:
                trPr = tr._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader"); trPr.append(th)
            for ci, txt in enumerate(row):
                cell = tr.cells[ci]
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
                tcW.set(qn("w:w"), str(tw[ci])); tcW.set(qn("w:type"), "dxa")
                if ri == 0:
                    b = OxmlElement("w:tcBorders"); bb = OxmlElement("w:bottom")
                    for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "0"), ("w:color", "4F81BD")):
                        bb.set(qn(k), v)
                    b.append(bb); tcPr.append(b)
                elif ri % 2 == 1:
                    shd = OxmlElement("w:shd")
                    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "D3DFEE")):
                        shd.set(qn(k), v)
                    tcPr.append(shd)
                par = cell.paragraphs[0]
                _spacing(par, 0, 0 if ri == 0 else 20, 1.0)
                add_runs(par, self.body_text(str(txt)), size, bold=(ri == 0) or (first_col_bold and ci == 0))
        cap = self.d.add_paragraph(); _spacing(cap, 60, 200)
        if caption:
            add_runs(cap, self.body_text(caption), 18, italic=True)

    def refs(self, entries):
        for n, text in entries:
            par = self.d.add_paragraph(); _spacing(par, 0, 60)
            add_runs(par, "[%d] %s" % (n, text), 18)

    def save(self, path):
        self.d.save(path)


# ------------------------------------------------------------------------------------------------ citations
class Citations:
    """Numbers citation keys in order of first appearance and renders [1], [2,3], [4–6]."""

    PAT = re.compile(r"\[@([^\]]+)\]")

    def __init__(self, library):
        self.lib = library
        self.order = []

    def scan(self, text):
        for m in self.PAT.finditer(text):
            for k in [x.strip().lstrip("@") for x in m.group(1).split(";")]:
                if k not in self.lib:
                    raise KeyError("unknown citation key " + k)
                if k not in self.order:
                    self.order.append(k)

    def num(self, k):
        return self.order.index(k) + 1

    def resolve(self, text):
        def rep(m):
            ns = sorted({self.num(x.strip().lstrip("@")) for x in m.group(1).split(";")})
            out, i = [], 0
            while i < len(ns):
                j = i
                while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
                    j += 1
                out.append(str(ns[i]) if j == i else ("%d,%d" % (ns[i], ns[j]) if j == i + 1 else "%d–%d" % (ns[i], ns[j])))
                i = j + 1
            return "[" + ",".join(out) + "]"
        return self.PAT.sub(rep, text)

    def entries(self):
        unused = [k for k in self.lib if k not in self.order]
        if unused:
            raise ValueError("references never cited: " + ", ".join(unused))
        return [(i + 1, self.lib[k]) for i, k in enumerate(self.order)]
