# -*- coding: utf-8 -*-
"""Cover letter for the Flow Measurement and Instrumentation submission."""
import os, json
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
R = json.load(open(os.path.join(HERE, "fmi_results.json")))
AB, SEL = R["ablation"], R["rejection_selectivity"]
UNC = AB["Uncorrected meter"]; ORC = AB["Correlation with known X_LM (oracle)"]
E2 = AB["E2 only (diagnostic)"]; FN = AB["Fusion, carryover omitted from E1 budget"]
FR = AB["Robust asymmetric fusion"]; SW = AB["SWCL: robust fusion + consistency rejection"]

DOI_PAPER = "10.5281/zenodo.22760861"
DOI_CODE = "10.5281/zenodo.22760863"
GITHUB = "https://github.com/sandlerleon/hybrid-liquid-loading-estimation"

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)


def p(t="", bold=False, italic=False, size=11, after=10):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(after)
    if t:
        r = par.add_run(t); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    return par


p("Leon Sandler", bold=True, after=0)
p("Independent Researcher", after=0)
p("Northbrook, IL 60062, USA", after=0)
p("sandler.leon@gmail.com", after=14)
p("15 September 2026", after=14)

p("The Editors", after=0)
p("Flow Measurement and Instrumentation", after=0)
p("Elsevier", after=14)

p("Dear Editors,", after=12)

p("I am submitting the enclosed manuscript, “Hybrid Physics–Diagnostic Estimation of Liquid "
  "Loading for Wet-Gas Flowmeter Correction: A Model-Based Uncertainty Framework,” for "
  "consideration in Flow Measurement and Instrumentation.")

p("The paper addresses a gap between two mature literatures. Wet-gas corrections for "
  "differential-pressure meters — Murdock, Chisholm, de Leeuw, ISO/TR 11583 — are well "
  "established, and all of them express the over-reading as a function of the Lockhart–Martinelli "
  "parameter. At a production metering station that parameter is not measured. The correlation is "
  "not the obstacle; its input is.")

p("Recent work in this journal addresses the same difficulty by learning the over-reading or the "
  "flow rates directly from experimental datasets, including the hybrid physics–machine-learning "
  "model of Alkhurayef and Alsarkhi (2026) and the cone-meter work of Pereira et al. (2026), "
  "while Nasr et al. (2026) resolve it by adding a third pressure tap. The present paper takes a "
  "different route: it treats liquid loading as a latent variable estimated from instrumentation "
  "already present, and leaves the established correlation unchanged. The learned component is "
  "reduced to a single scalar whose uncertainty can be stated and propagated, and because a "
  "second, structurally distinct estimate of that scalar exists, the two can be cross-checked — "
  "which is what makes it possible for the system to withhold a correction it cannot justify.")

p("Three findings may be of particular interest to your readers, and I would highlight that the "
  "first two are negative results about the method as initially formulated rather than "
  "confirmations of it.", bold=True)

p(f"First, combining the two estimators by inverse-variance weighting is worse "
  f"(±{FN['band95']:.2f}%) than using the diagnostic estimator alone (±{E2['band95']:.2f}%) "
  f"unless the thermodynamic estimator declares an uncertainty that includes an allowance for the "
  f"upstream carryover it cannot observe. This is a general point about fusing estimators with "
  f"structurally different failure modes: inverse-variance weighting is only as good as the "
  f"declared variances. With that allowance and an asymmetric weighting motivated by the "
  f"one-directional nature of the failure, the simulated band improves to ±{FR['band95']:.2f}%.")

p(f"Second, the two estimates are not statistically independent, and treating them as though they "
  f"were does not degrade the acceptance test gracefully. It progressively disables it: as the "
  f"correlation between the estimator errors rises, the consistency index shrinks and the test "
  f"stops firing, accepting everything. For a mechanism whose purpose is to withhold a "
  f"correction, that is the worst available failure mode. The covariance term is therefore "
  f"retained explicitly, and the paper recommends both excluding the thermodynamic estimate from "
  f"the diagnostic feature vector and measuring the residual correlation during calibration.")

p(f"Third, the consistency index is a specific detector rather than a blunt filter: upstream "
  f"carryover is present in {100*SEL['all']:.0f}% of simulated realisations but in "
  f"{100*SEL['rejected']:.0f}% of those the test rejects.")

p("I want to be explicit with the editors about the status of the evidence, because it determines "
  "how the paper should be reviewed. This is a computational study. No wet-gas loop data were "
  "used, the diagnostic estimator is specified in enough detail to be fitted and reproduced but "
  "is not fitted here, and every performance figure is a modelled uncertainty conditional on "
  "tabulated assumptions rather than a measured accuracy. The title, abstract, results and "
  "conclusions all say so. The experiment that would test the framework — facility, test matrix, "
  "the four comparisons, and the criteria that would refute it — is specified in the Methods "
  "rather than deferred to the discussion, so that the claims can be read against the test that "
  "would falsify them.", bold=True)

p(f"Because the entire quantitative content is computational, the complete reproducibility "
  f"package is deposited openly rather than promised on request. The Python implementation, the "
  f"simulation configuration, the fixed random seeds, the figure generators and the generated "
  f"summary data are available at {GITHUB} and archived at https://doi.org/{DOI_CODE}. Running "
  f"the deposited scripts reproduces every table and figure in the manuscript; the manuscript "
  f"builder itself reads its numbers from the generated JSON summaries, so the text and the "
  f"computation cannot diverge. A preprint is archived at https://doi.org/{DOI_PAPER}.")

p("The manuscript is original, is not under consideration elsewhere, and has not been published "
  "previously except as the archived preprint noted above. I am the sole author, I have no "
  "competing interests, and the work received no external funding. Use of a generative AI tool as "
  "a drafting and literature-synthesis aid is disclosed in the manuscript.")

p("Thank you for considering this submission.", after=16)

p("Sincerely,", after=16)
p("Leon Sandler", bold=True, after=0)
p("Independent Researcher, Northbrook, IL, USA", after=0)
p("sandler.leon@gmail.com", after=0)

out = os.path.join(HERE, "FMI_Cover_Letter_Sandler.docx")
doc.save(out)
print("saved", out)
